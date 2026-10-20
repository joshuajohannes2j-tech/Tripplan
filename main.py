import streamlit as st
import json
import os
import time as time_module
from datetime import time, date, timedelta
from io import BytesIO

import requests
import pandas as pd
import folium
from streamlit_folium import st_folium

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter

import plotly.express as px


st.set_page_config(
    page_title="TripPlan | Your Next Adventure",
    page_icon="✈️",
    layout="wide",
    initial_sidebar_state="collapsed"
)


# --- TripPlan Permanent Storage ---
DATA_FILE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "tripplan_data.json"
)


def save_trip_data():
    """Save trips and wishlist to a local JSON file."""
    data = {
        "trips": st.session_state.get("trips", []),
        "wishlist": st.session_state.get("wishlist", [])
    }

    temporary_file = DATA_FILE + ".tmp"

    with open(temporary_file, "w", encoding="utf-8") as file:
        json.dump(data, file, indent=4, ensure_ascii=False)

    os.replace(temporary_file, DATA_FILE)


def load_trip_data():
    """Load trips and wishlist from the local JSON file."""
    if not os.path.exists(DATA_FILE):
        return {"trips": [], "wishlist": []}

    try:
        with open(DATA_FILE, "r", encoding="utf-8") as file:
            data = json.load(file)

        if not isinstance(data, dict):
            raise ValueError("Invalid TripPlan data format.")

        trips = data.get("trips", [])
        wishlist = data.get("wishlist", [])

        if not isinstance(trips, list) or not isinstance(wishlist, list):
            raise ValueError("Trips and wishlist must be lists.")

        return {
            "trips": trips,
            "wishlist": wishlist
        }

    except (json.JSONDecodeError, OSError, ValueError) as error:
        st.error(
            "TripPlan couldn't load its saved data. "
            "Your data file has been left unchanged."
        )
        st.stop()


# Load saved data once per Streamlit session
if "trip_data_loaded" not in st.session_state:
    saved_data = load_trip_data()

    st.session_state.trips = saved_data["trips"]
    st.session_state.wishlist = saved_data["wishlist"]
    st.session_state.trip_data_loaded = True

# --- End TripPlan Permanent Storage ---


# --- Excel Export Helpers ---


def create_tripplan_excel():
    trips = st.session_state.get("trips", [])
    wishlist = st.session_state.get("wishlist", [])

    output = BytesIO()

    # ----------------------------------
    # PREPARE TRIP DATA
    # ----------------------------------
    trip_rows = []
    detail_rows = []

    total_expenses = 0.0
    total_budget = 0.0
    currencies = set()

    for trip_index, trip in enumerate(trips, start=1):
        if not isinstance(trip, dict):
            continue

        destination = trip.get("destination", "Unknown destination")
        start_date = trip.get("start_date", "")
        end_date = trip.get("end_date", "")
        currency = trip.get("currency", "ZAR")
        budget = trip.get("budget", 0)

        try:
            budget = float(budget or 0)
        except (TypeError, ValueError):
            budget = 0.0

        currencies.add(currency)
        total_budget += budget

        expenses = trip.get("expenses", [])
        if not isinstance(expenses, list):
            expenses = []

        trip_spent = 0.0

        for expense in expenses:
            if not isinstance(expense, dict):
                continue

            amount = expense.get("amount", 0)

            try:
                amount = float(amount or 0)
            except (TypeError, ValueError):
                amount = 0.0

            expense_currency = expense.get("currency", currency)
            category = expense.get("category", "Other")
            description = expense.get(
                "name",
                expense.get("description", "")
            )
            notes = expense.get("notes", "")

            trip_spent += amount

            detail_rows.append({
                "Type": "Expense",
                "Trip": destination,
                "Date": expense.get("date", ""),
                "Name / Description": description,
                "Category": category,
                "Amount": amount,
                "Currency": expense_currency,
                "Notes": notes,
            })

        total_expenses += trip_spent

        trip_rows.append({
            "Trip #": trip_index,
            "Destination": destination,
            "Start Date": start_date,
            "End Date": end_date,
            "Currency": currency,
            "Budget": budget,
            "Spent": trip_spent,
            "Remaining": budget - trip_spent,
        })

        # ----------------------------------
        # EXPORT ITINERARY
        # ----------------------------------
        itinerary = trip.get("itinerary", [])

        if isinstance(itinerary, list):
            for item in itinerary:
                if isinstance(item, dict):
                    detail_rows.append({
                        "Type": "Itinerary",
                        "Trip": destination,
                        "Date": item.get("date", ""),
                        "Name / Description": item.get(
                            "activity",
                            item.get("name", "")
                        ),
                        "Category": "Activity",
                        "Amount": "",
                        "Currency": currency,
                        "Notes": item.get("notes", ""),
                    })

                elif isinstance(item, str):
                    detail_rows.append({
                        "Type": "Itinerary",
                        "Trip": destination,
                        "Date": "",
                        "Name / Description": item,
                        "Category": "Activity",
                        "Amount": "",
                        "Currency": currency,
                        "Notes": "",
                    })

    # ----------------------------------
    # EXPORT WISHLIST
    # ----------------------------------
    for item in wishlist:
        if isinstance(item, dict):
            detail_rows.append({
                "Type": "Wishlist",
                "Trip": "",
                "Date": "",
                "Name / Description": item.get(
                    "name",
                    item.get("destination", str(item))
                ),
                "Category": item.get("category", "Destination"),
                "Amount": "",
                "Currency": "",
                "Notes": item.get("notes", ""),
            })

        else:
            detail_rows.append({
                "Type": "Wishlist",
                "Trip": "",
                "Date": "",
                "Name / Description": str(item),
                "Category": "Destination",
                "Amount": "",
                "Currency": "",
                "Notes": "",
            })

    # ----------------------------------
    # CREATE DATAFRAMES
    # ----------------------------------
    trip_columns = [
        "Trip #", "Destination", "Start Date", "End Date",
        "Currency", "Budget", "Spent", "Remaining"
    ]

    detail_columns = [
        "Type", "Trip", "Date", "Name / Description",
        "Category", "Amount", "Currency", "Notes"
    ]

    trips_df = pd.DataFrame(trip_rows, columns=trip_columns)
    details_df = pd.DataFrame(detail_rows, columns=detail_columns)

    if trips_df.empty:
        trips_df = pd.DataFrame([{
            "Trip #": "",
            "Destination": "No trips added yet",
            "Start Date": "",
            "End Date": "",
            "Currency": "",
            "Budget": 0,
            "Spent": 0,
            "Remaining": 0,
        }], columns=trip_columns)

    if details_df.empty:
        details_df = pd.DataFrame([{
            "Type": "",
            "Trip": "",
            "Date": "",
            "Name / Description": "No expenses, itinerary items or wishlist entries yet",
            "Category": "",
            "Amount": "",
            "Currency": "",
            "Notes": "",
        }], columns=detail_columns)

    # ----------------------------------
    # CREATE WORKBOOK WITH TWO SHEETS
    # ----------------------------------
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        trips_df.to_excel(
            writer,
            sheet_name="Overview",
            index=False,
            startrow=7
        )

        details_df.to_excel(
            writer,
            sheet_name="All Details",
            index=False,
            startrow=3
        )

        ws = writer.sheets["Overview"]
        detail_ws = writer.sheets["All Details"]

        # ----------------------------------
        # COLOURS AND STYLES
        # ----------------------------------
        navy = "14243A"
        teal = "16B8A6"
        white = "FFFFFF"
        light_teal = "E8F8F5"
        light_gray = "F3F6FA"
        dark_text = "243447"

        thin_gray = Side(style="thin", color="D9E2EC")

        title_font = Font(
            name="Aptos Display",
            size=22,
            bold=True,
            color=white
        )

        subtitle_font = Font(
            name="Aptos",
            size=10,
            color="DCE6F1"
        )

        header_font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=white
        )

        body_font = Font(
            name="Aptos",
            size=10,
            color=dark_text
        )

        # ----------------------------------
        # OVERVIEW TITLE
        # ----------------------------------
        ws.merge_cells("A1:H2")
        ws["A1"] = "TRIPPLAN | TRAVEL REPORT"
        ws["A1"].font = title_font
        ws["A1"].fill = PatternFill("solid", fgColor=navy)
        ws["A1"].alignment = Alignment(
            horizontal="left",
            vertical="center"
        )

        for row in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=8):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=navy)

        ws.row_dimensions[1].height = 25
        ws.row_dimensions[2].height = 18

        ws.merge_cells("A3:H3")
        ws["A3"] = "Your trips, budgets and spending at a glance"
        ws["A3"].font = subtitle_font
        ws["A3"].fill = PatternFill("solid", fgColor=navy)
        ws["A3"].alignment = Alignment(vertical="center")

        for cell in ws[3][:8]:
            cell.fill = PatternFill("solid", fgColor=navy)

        ws.row_dimensions[3].height = 24

        # ----------------------------------
        # SUMMARY CARDS
        # ----------------------------------
        trip_count = len(trip_rows)

        # Avoid combining different currencies into one misleading total.
        single_currency = len(currencies) <= 1
        currency_label = next(iter(currencies), "ZAR") if single_currency else "Mixed currencies"

        summary_items = [
            ("A5", "TOTAL TRIPS", trip_count),
            ("C5", "TOTAL BUDGET", total_budget if single_currency else "See trips"),
            ("E5", "TOTAL SPENT", total_expenses if single_currency else "See expenses"),
            (
                "G5",
                "BUDGET LEFT",
                total_budget - total_expenses if single_currency else "See trips"
            ),
        ]

        for cell_ref, label, value in summary_items:
            cell = ws[cell_ref]
            cell.value = label
            cell.font = Font(bold=True, size=9, color=teal)
            cell.fill = PatternFill("solid", fgColor=light_teal)
            cell.alignment = Alignment(vertical="center")

            value_cell = ws.cell(
                row=6,
                column=cell.column,
                value=value
            )
            value_cell.font = Font(
                size=15,
                bold=True,
                color=navy
            )
            value_cell.fill = PatternFill("solid", fgColor=light_teal)
            value_cell.alignment = Alignment(vertical="center")

            if isinstance(value, (int, float)) and cell.column != 1:
                value_cell.number_format = '#,##0.00'

        ws.merge_cells("A7:H7")
        ws["A7"] = (
            f"Currency: {currency_label} | "
            "Individual trip amounts remain in their recorded currencies."
        )
        ws["A7"].font = Font(size=9, italic=True, color=dark_text)

        # ----------------------------------
        # FORMAT OVERVIEW TABLE
        # ----------------------------------
        header_row = 8

        for cell in ws[header_row]:
            if cell.column <= len(trip_columns):
                cell.fill = PatternFill("solid", fgColor=navy)
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )
                cell.border = Border(bottom=Side(style="medium", color=teal))

        ws.row_dimensions[header_row].height = 28
        ws.freeze_panes = "A9"
        ws.auto_filter.ref = (
            f"A{header_row}:H{header_row + len(trips_df)}"
        )

        for row in ws.iter_rows(
            min_row=header_row + 1,
            max_row=header_row + len(trips_df),
            min_col=1,
            max_col=8
        ):
            for cell in row:
                cell.font = body_font
                cell.border = Border(bottom=thin_gray)

                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=light_gray)

            for column in (6, 7, 8):
                ws.cell(row=row[0].row, column=column).number_format = '#,##0.00'

        # ----------------------------------
        # FORMAT ALL DETAILS SHEET
        # ----------------------------------
        detail_ws.merge_cells("A1:H1")
        detail_ws["A1"] = "TRIPPLAN | ALL DETAILS"
        detail_ws["A1"].font = Font(
            name="Aptos Display",
            size=20,
            bold=True,
            color=white
        )
        detail_ws["A1"].fill = PatternFill("solid", fgColor=navy)
        detail_ws["A1"].alignment = Alignment(vertical="center")

        for row in detail_ws.iter_rows(min_row=1, max_row=1, min_col=1, max_col=8):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=navy)

        detail_ws.row_dimensions[1].height = 36

        detail_ws.merge_cells("A2:H2")
        detail_ws["A2"] = "Expenses, itinerary activities and wishlist destinations"
        detail_ws["A2"].font = Font(size=10, color=dark_text, italic=True)

        detail_header_row = 4

        for cell in detail_ws[detail_header_row]:
            if cell.column <= len(detail_columns):
                cell.fill = PatternFill("solid", fgColor=navy)
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )
                cell.border = Border(bottom=Side(style="medium", color=teal))

        detail_ws.row_dimensions[detail_header_row].height = 28
        detail_ws.freeze_panes = "A5"
        detail_ws.auto_filter.ref = (
            f"A{detail_header_row}:H{detail_header_row + len(details_df)}"
        )

        for row in detail_ws.iter_rows(
            min_row=detail_header_row + 1,
            max_row=detail_header_row + len(details_df),
            min_col=1,
            max_col=8
        ):
            for cell in row:
                cell.font = body_font
                cell.border = Border(bottom=thin_gray)
                cell.alignment = Alignment(vertical="top", wrap_text=True)

                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=light_gray)

            amount_cell = detail_ws.cell(row=row[0].row, column=6)
            if isinstance(amount_cell.value, (int, float)):
                amount_cell.number_format = '#,##0.00'

        # ----------------------------------
        # COLUMN WIDTHS
        # ----------------------------------
        overview_widths = {
            "A": 12,
            "B": 24,
            "C": 16,
            "D": 16,
            "E": 16,
            "F": 16,
            "G": 16,
            "H": 16,
        }

        for column, width in overview_widths.items():
            ws.column_dimensions[column].width = width

        detail_widths = {
            "A": 16,
            "B": 22,
            "C": 16,
            "D": 38,
            "E": 20,
            "F": 16,
            "G": 16,
            "H": 38,
        }

        for column, width in detail_widths.items():
            detail_ws.column_dimensions[column].width = width

        # ----------------------------------
        # PRINT SETTINGS
        # ----------------------------------
        for sheet in (ws, detail_ws):
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.page_margins.left = 0.3
            sheet.page_margins.right = 0.3
            sheet.page_margins.top = 0.4
            sheet.page_margins.bottom = 0.4
            sheet.oddFooter.center.text = "TripPlan Travel Report"
            sheet.oddFooter.right.text = "Page &P of &N"

    output.seek(0)
    return output.getvalue()




# --- TripPlan Aurora Startup Screen ---
if "startup_screen_shown" not in st.session_state:
    st.session_state.startup_screen_shown = False

if not st.session_state.startup_screen_shown:
    startup_placeholder = st.empty()

    with startup_placeholder:
        st.components.v1.html(
            """
            <!DOCTYPE html>
            <html>
            <head>
            <meta name="viewport" content="width=device-width, initial-scale=1.0">
            <style>
                * {
                    box-sizing: border-box;
                }

                html, body {
                    margin: 0;
                    padding: 0;
                    width: 100%;
                    height: 620px;
                    overflow: hidden;
                    font-family: Arial, sans-serif;
                }

                .aurora {
                    height: 620px;
                    width: 100%;
                    display: flex;
                    flex-direction: column;
                    align-items: center;
                    justify-content: center;
                    color: white;
                    text-align: center;
                    overflow: hidden;
                    background:
                        radial-gradient(
                            circle at 50% 42%,
                            rgba(0, 126, 255, 0.34),
                            transparent 36%
                        ),
                        radial-gradient(
                            circle at 15% 15%,
                            rgba(67, 77, 255, 0.18),
                            transparent 35%
                        ),
                        linear-gradient(
                            145deg,
                            #030b1c,
                            #071f49 55%,
                            #020817
                        );
                }

                .globe-wrap {
                    width: 180px;
                    height: 180px;
                    display: flex;
                    align-items: center;
                    justify-content: center;
                    position: relative;
                    margin-bottom: 25px;
                }

                .glow {
                    position: absolute;
                    width: 150px;
                    height: 150px;
                    border-radius: 50%;
                    background: #087cff;
                    filter: blur(38px);
                    opacity: 0.6;
                    animation: pulse 2s ease-in-out infinite;
                }

                .globe {
                    position: relative;
                    width: 142px;
                    height: 142px;
                    overflow: hidden;
                    border-radius: 50%;
                    border: 2px solid #75d5ff;
                    background:
                        radial-gradient(
                            circle at 30% 25%,
                            #a4edff,
                            transparent 24%
                        ),
                        radial-gradient(
                            circle at 70% 70%,
                            #0448bd,
                            transparent 55%
                        ),
                        linear-gradient(
                            135deg,
                            #44ceff,
                            #087cff 55%,
                            #061c52
                        );
                    box-shadow:
                        0 0 25px #087cff,
                        0 0 65px rgba(8, 124, 255, 0.65),
                        inset -15px -12px 25px rgba(0, 0, 30, 0.6);
                    animation: float 3s ease-in-out infinite;
                }

                .globe::before {
                    content: "";
                    position: absolute;
                    inset: -20%;
                    border-radius: 50%;
                    background: repeating-linear-gradient(
                        90deg,
                        transparent 0 24%,
                        rgba(210, 245, 255, 0.6) 25%,
                        transparent 26% 49%
                    );
                    animation: rotate 8s linear infinite;
                }

                .orbit {
                    position: absolute;
                    width: 205px;
                    height: 65px;
                    border: 1px solid rgba(115, 210, 255, 0.85);
                    border-radius: 50%;
                    transform: rotate(-25deg);
                    box-shadow: 0 0 12px rgba(8, 124, 255, 0.55);
                }

                .brand {
                    margin: 0;
                    font-size: clamp(36px, 5vw, 48px);
                    font-weight: 800;
                    letter-spacing: -1.5px;
                    text-shadow: 0 0 24px rgba(98, 199, 255, 0.35);
                }

                .brand span {
                    color: #62c7ff;
                }

                .tagline {
                    margin: 12px 15px 0;
                    color: #c4d9f8;
                    font-size: 15px;
                    letter-spacing: 1px;
                }

                .track {
                    width: 180px;
                    height: 4px;
                    margin-top: 35px;
                    overflow: hidden;
                    border-radius: 10px;
                    background: rgba(180, 210, 255, 0.2);
                }

                .bar {
                    height: 100%;
                    width: 42%;
                    background: linear-gradient(
                        90deg,
                        #087cff,
                        #75dcff,
                        white
                    );
                    box-shadow: 0 0 12px #35a4f5;
                    animation: loading 1.3s ease-in-out infinite;
                }

                .loading-text {
                    margin-top: 13px;
                    color: #91b9e9;
                    font-size: 11px;
                    letter-spacing: 3px;
                    text-transform: uppercase;
                }

                @keyframes pulse {
                    0%, 100% { opacity: 0.35; transform: scale(0.9); }
                    50% { opacity: 0.75; transform: scale(1.08); }
                }

                @keyframes float {
                    0%, 100% { transform: translateY(0); }
                    50% { transform: translateY(-7px); }
                }

                @keyframes rotate {
                    from { transform: rotate(0deg); }
                    to { transform: rotate(360deg); }
                }

                @keyframes loading {
                    from { transform: translateX(-110%); }
                    to { transform: translateX(250%); }
                }

                @media (prefers-reduced-motion: reduce) {
                    *, *::before, *::after {
                        animation: none !important;
                    }
                }
            </style>
            </head>
            <body>
                <div class="aurora">
                    <div class="globe-wrap">
                        <div class="glow"></div>
                        <div class="globe"></div>
                        <div class="orbit"></div>
                    </div>

                    <h1 class="brand">Trip<span>Plan</span></h1>

                    <p class="tagline">
                        Your next adventure starts here.
                    </p>

                    <div class="track">
                        <div class="bar"></div>
                    </div>

                    <div class="loading-text">
                        Preparing your journey
                    </div>
                </div>
            </body>
            </html>
            """,
            height=620,
            scrolling=False
        )

    time_module.sleep(5)
    startup_placeholder.empty()
    st.session_state.startup_screen_shown = True

# --- End TripPlan Aurora Startup Screen ---

@st.cache_data(ttl=3600)
def search_real_places(query):
    """Search for real places using OpenStreetMap's Nominatim service."""
    response = requests.get(
        "https://nominatim.openstreetmap.org/search",
        params={
            "q": query,
            "format": "jsonv2",
            "addressdetails": 1,
            "limit": 5,
        },
        headers={"User-Agent": "TripPlanTravelPlanner/1.0"},
        timeout=10,
    )
    response.raise_for_status()
    return response.json()


@st.cache_data(ttl=86400)
def get_destination_coordinates(destination):
    url = "https://geocoding-api.open-meteo.com/v1/search"

    response = requests.get(
        url,
        params={
            "name": destination,
            "count": 1,
            "language": "en",
            "format": "json"
        },
        timeout=10
    )

    response.raise_for_status()
    data = response.json()

    if not data.get("results"):
        return None

    place = data["results"][0]

    return {
        "latitude": place["latitude"],
        "longitude": place["longitude"],
        "name": place["name"],
        "country": place.get("country", "")
    }

def get_weather(destination):
    # Find the destination's coordinates
    geo_url = "https://geocoding-api.open-meteo.com/v1/search"

    geo_response = requests.get(
        geo_url,
        params={
            "name": destination,
            "count": 1,
            "language": "en",
            "format": "json"
        },
        timeout=10
    )
    geo_response.raise_for_status()
    geo_data = geo_response.json()

    if not geo_data.get("results"):
        return None

    place = geo_data["results"][0]
    latitude = place["latitude"]
    longitude = place["longitude"]

    # Retrieve current conditions and a 7-day forecast
    weather_url = "https://api.open-meteo.com/v1/forecast"

    weather_response = requests.get(
        weather_url,
        params={
            "latitude": latitude,
            "longitude": longitude,
            "current": (
                "temperature_2m,relative_humidity_2m,"
                "apparent_temperature,precipitation,"
                "weather_code,wind_speed_10m"
            ),
            "daily": (
                "weather_code,temperature_2m_max,"
                "temperature_2m_min,precipitation_probability_max"
            ),
            "timezone": "auto",
            "forecast_days": 7
        },
        timeout=10
    )
    weather_response.raise_for_status()

    return {
        "place": place,
        "weather": weather_response.json()
    }


def weather_description(code):
    descriptions = {
        0: "Clear sky",
        1: "Mainly clear",
        2: "Partly cloudy",
        3: "Overcast",
        45: "Foggy",
        48: "Rime fog",
        51: "Light drizzle",
        53: "Drizzle",
        55: "Heavy drizzle",
        61: "Light rain",
        63: "Rain",
        65: "Heavy rain",
        71: "Light snow",
        73: "Snow",
        75: "Heavy snow",
        80: "Rain showers",
        81: "Moderate rain showers",
        82: "Heavy rain showers",
        95: "Thunderstorm",
        96: "Thunderstorm with hail",
        99: "Severe thunderstorm"
    }
    return descriptions.get(code, "Conditions unavailable")

# --- Page navigation setup ---
if "page" not in st.session_state:
    st.session_state.page = "Home"


def go_to(page):
    st.session_state.page = page

    if page == "Dashboard":
        if "trips" not in st.session_state:
            st.session_state.trips = []

        if "wishlist" not in st.session_state:
            st.session_state.wishlist = []


# --- Styling ---
st.markdown("""
<style>
    /* Main background */
    .stApp {
        background: linear-gradient(180deg, #f4f9ff 0%, #ffffff 75%);
    }

    /* Main content width and spacing */
        .block-container {
        padding-top: 5rem;
        padding-bottom: 3rem;
        max-width: 1250px;
    }

    /* Headings */
    h1, h2, h3 {
        color: #12366b;
        letter-spacing: -0.5px;
    }

    p {
        color: #536b89;
        line-height: 1.8;
        font-size: 16px;
    }

    /* All buttons */
    div.stButton > button {
        background: #087cff !important;
        color: #ffffff !important;
        border: 1px solid #087cff !important;
        border-radius: 12px !important;
        min-height: 46px;
        padding: 0.65rem 1.1rem;
        font-weight: 600 !important;
        box-shadow: 0 4px 12px rgba(8, 124, 255, 0.12);
        transition: all 0.2s ease;
    }

    div.stButton > button p {
        color: #ffffff !important;
        font-size: 15px !important;
        font-weight: 600 !important;
    }

    div.stButton > button:hover {
        background: #0563d1 !important;
        border-color: #0563d1 !important;
        box-shadow: 0 6px 18px rgba(8, 124, 255, 0.22);
        transform: translateY(-1px);
    }

    div.stButton > button:focus {
        color: #ffffff !important;
        border-color: #087cff !important;
    }

    /* Divider */
    hr {
        border-color: #dce9f8;
        margin-top: 1rem;
        margin-bottom: 1.8rem;
    }

    /* Images */
    [data-testid="stImage"] img {
        border-radius: 22px;
    }

    /* Responsive spacing */
    @media (max-width: 768px) {
        .block-container {
            padding-top: 1rem;
            padding-left: 1rem;
            padding-right: 1rem;
        }
    }
    

    /* Dashboard welcome banner */
    .dashboard-hero {
        position: relative;
        overflow: hidden;
        background:
            radial-gradient(
                circle at 88% 15%,
                rgba(255, 255, 255, 0.22),
                transparent 30%
            ),
            linear-gradient(
                120deg,
                #073b87 0%,
                #087cff 58%,
                #35a4f5 100%
            );
        padding: 44px 42px;
        border-radius: 26px;
        color: #ffffff;
        margin: 18px 0 30px 0;
        box-shadow: 0 16px 38px rgba(8, 92, 190, 0.20);
        border: 1px solid rgba(255, 255, 255, 0.25);
    }

    .dashboard-hero .hero-label {
        display: inline-block;
        color: #ffffff;
        background: rgba(255, 255, 255, 0.14);
        border: 1px solid rgba(255, 255, 255, 0.30);
        padding: 8px 13px;
        border-radius: 999px;
        font-size: 11px;
        font-weight: 700;
        letter-spacing: 1.3px;
    }

    .dashboard-hero h1 {
        color: #ffffff !important;
        font-size: clamp(30px, 4vw, 44px);
        line-height: 1.16;
        letter-spacing: -1.2px;
        margin-top: 20px;
        margin-bottom: 14px;
    }

    .dashboard-hero p {
        color: #eaf4ff !important;
        font-size: 16px;
        line-height: 1.8;
        max-width: 680px;
        margin-bottom: 0;
    }

    /* Make the welcome banner comfortable on smaller screens */
    @media (max-width: 768px) {
        .dashboard-hero {
            padding: 30px 24px;
            border-radius: 20px;
            margin-top: 12px;
        }

        .dashboard-hero h1 {
            font-size: 30px;
            letter-spacing: -0.6px;
        }

        .dashboard-hero p {
            font-size: 15px;
            line-height: 1.7;
        }

        .dashboard-hero .hero-label {
            font-size: 10px;
            letter-spacing: 0.8px;
        }
    }

    
    
    /* Premium dashboard overview cards */
    div[data-testid="stMetric"] {
        background: linear-gradient(
            145deg,
            #ffffff 0%,
            #f1f7ff 100%
        );
        padding: 26px 24px;
        border: 1px solid #d8e7fb;
        border-radius: 22px;
        box-shadow: 0 8px 24px rgba(26, 71, 125, 0.07);
        transition: transform 0.25s ease,
                    box-shadow 0.25s ease,
                    border-color 0.25s ease;
        min-height: 145px;
    }

    div[data-testid="stMetric"]:hover {
        transform: translateY(-5px);
        box-shadow: 0 16px 32px rgba(26, 71, 125, 0.13);
        border-color: #9fc8ff;
    }

    div[data-testid="stMetricLabel"] {
        color: #607794;
        font-size: 14px;
        font-weight: 700;
        margin-bottom: 10px;
    }

    div[data-testid="stMetricValue"] {
        color: #087cff;
        font-size: 34px;
        font-weight: 800;
        letter-spacing: -0.8px;
    }

    /* Smaller screens */
    @media (max-width: 768px) {
        div[data-testid="stMetric"] {
            padding: 20px 16px;
            border-radius: 18px;
            min-height: 120px;
        }

        div[data-testid="stMetricValue"] {
            font-size: 28px;
        }
    }
    
    /* Dashboard section headings */
    .dashboard-section {
        color: #12366b;
        font-size: 23px;
        font-weight: 700;
        margin-top: 12px;
        margin-bottom: 12px;
    }
    
    
    /* ==========================================
       TRIPPLAN — AURORA VISUAL UPGRADE
       Visual styling only — no app logic changed
       ========================================== */

    /* Soft atmospheric background */
    .stApp {
        background:
            radial-gradient(
                ellipse at 5% 5%,
                rgba(8, 124, 255, 0.09),
                transparent 32%
            ),
            radial-gradient(
                ellipse at 95% 12%,
                rgba(53, 164, 245, 0.10),
                transparent 30%
            ),
            linear-gradient(
                180deg,
                #f2f7ff 0%,
                #ffffff 58%,
                #f7faff 100%
            );
    }

    /* Keep your preferred top spacing */
    .block-container {
        padding-top: 5rem;
        padding-bottom: 4rem;
        max-width: 1250px;
    }

    /* More polished section headings */
    .dashboard-section {
        position: relative;
        color: #12366b;
        font-size: 24px;
        font-weight: 800;
        letter-spacing: -0.6px;
        margin-top: 1.6rem;
        margin-bottom: 1.2rem;
        padding-bottom: 12px;
        border-bottom: 1px solid #e1ebf9;
    }

    .dashboard-section::after {
        content: "";
        position: absolute;
        bottom: -1px;
        left: 0;
        width: 64px;
        height: 3px;
        border-radius: 999px;
        background: linear-gradient(
            90deg,
            #087cff,
            #48b8ff
        );
    }

    /* Give the overview cards a premium finish */
    div[data-testid="stMetric"] {
        position: relative;
        overflow: hidden;
        background: linear-gradient(
            145deg,
            #ffffff 0%,
            #f3f8ff 100%
        );
        border: 1px solid #dce8f8;
        border-radius: 22px;
        box-shadow: 0 10px 28px rgba(26, 71, 125, 0.07);
        transition:
            transform 0.25s ease,
            box-shadow 0.25s ease,
            border-color 0.25s ease;
    }

    /* Thin blue accent along each card's top edge */
    div[data-testid="stMetric"]::before {
        content: "";
        position: absolute;
        top: 0;
        left: 0;
        right: 0;
        height: 4px;
        background: linear-gradient(
            90deg,
            #0758ca,
            #087cff,
            #48b8ff
        );
        opacity: 0.8;
    }

    /* Smooth hover interaction */
    div[data-testid="stMetric"]:hover {
        transform: translateY(-4px);
        border-color: #a9ceff;
        box-shadow: 0 17px 36px rgba(26, 71, 125, 0.13);
    }

    div[data-testid="stMetricValue"] {
        color: #0758ca;
        font-weight: 800;
        letter-spacing: -0.8px;
    }

    div[data-testid="stMetricLabel"] {
        color: #526b8c;
        font-weight: 650;
    }

    /* Cleaner form fields throughout the app */
    div[data-testid="stTextInput"] input,
    div[data-testid="stTextArea"] textarea {
        border-radius: 12px;
        border-color: #d4e2f5;
        background-color: #ffffff;
        transition:
            border-color 0.2s ease,
            box-shadow 0.2s ease;
    }

    div[data-testid="stTextInput"] input:focus,
    div[data-testid="stTextArea"] textarea:focus {
        border-color: #087cff;
        box-shadow: 0 0 0 3px rgba(8, 124, 255, 0.12);
    }

    /* Responsive refinements */
    @media (max-width: 768px) {
        .dashboard-section {
            font-size: 21px;
            margin-top: 1.3rem;
        }

        div[data-testid="stMetric"] {
            border-radius: 18px;
        }

        .block-container {
            padding-bottom: 2.5rem;
        }
    }
    
    
    
    /* ==========================================
       TRAVEL MAP — PREMIUM FINISH
       ========================================== */

    /* Round and elevate the interactive map */
    div[data-testid="stCustomComponentV1"] {
        border-radius: 20px;
        overflow: hidden;
        box-shadow: 0 12px 30px rgba(26, 71, 125, 0.10);
    }

    /* Give dashboard section headings a little more polish */
    h2 {
        color: #12366b !important;
        font-weight: 800 !important;
        letter-spacing: -0.7px;
    }

    /* Style the map filter controls */
    div[data-testid="stRadio"] {
        background: rgba(255, 255, 255, 0.78);
        border: 1px solid #dce8f8;
        border-radius: 16px;
        padding: 12px 16px;
        margin: 10px 0 18px 0;
    }

    div[data-testid="stRadio"] input[type="radio"] {
        accent-color: #087cff;
    }

    /* ==========================================
       DESTINATION WEATHER — PREMIUM FINISH
       ========================================== */

    /* Rounded destination selector */
    div[data-testid="stSelectbox"] div[data-baseweb="select"] > div {
        border-radius: 12px;
        border-color: #d4e2f5;
        background-color: #ffffff;
        transition: border-color 0.2s ease,
                    box-shadow 0.2s ease;
    }

    div[data-testid="stSelectbox"] div[data-baseweb="select"] > div:hover {
        border-color: #87baff;
    }

    /* Make the forecast table feel like a polished panel */
    div[data-testid="stDataFrame"] {
        border: 1px solid #dce8f8;
        border-radius: 16px;
        overflow: hidden;
        box-shadow: 0 8px 24px rgba(26, 71, 125, 0.07);
    }

    /* Smooth, rounded feedback messages */
    div[data-testid="stAlert"] {
        border-radius: 14px;
    }

    /* Consistent section spacing */
    div[data-testid="stVerticalBlock"] > div {
        scroll-margin-top: 24px;
    }

    /* Smaller screens */
    @media (max-width: 768px) {
        div[data-testid="stRadio"] {
            padding: 10px;
            border-radius: 12px;
        }

        div[data-testid="stCustomComponentV1"] {
            border-radius: 14px;
        }

        div[data-testid="stDataFrame"] {
            border-radius: 12px;
        }
    }
    
    /* --- TripPlan Aurora Dashboard Upgrade --- */

.dashboard-hero {
    position: relative;
    overflow: hidden;
    isolation: isolate;
    padding: clamp(30px, 5vw, 52px);
    border: 1px solid rgba(132, 200, 255, 0.35);
    border-radius: 26px;
    background:
        radial-gradient(
            circle at 85% 15%,
            rgba(70, 186, 255, 0.35),
            transparent 35%
        ),
        linear-gradient(
            125deg,
            #061636 0%,
            #084ba5 55%,
            #087cff 100%
        );
    box-shadow: 0 14px 38px rgba(8, 75, 165, 0.2);
    animation: dashboardFadeUp 0.7s ease-out both;
}

.dashboard-hero::after {
    content: "";
    position: absolute;
    z-index: -1;
    right: -65px;
    top: -95px;
    width: 250px;
    height: 250px;
    border-radius: 50%;
    background: rgba(98, 199, 255, 0.28);
    filter: blur(45px);
    pointer-events: none;
}

.dashboard-hero h1 {
    position: relative;
    color: #ffffff !important;
    font-size: clamp(28px, 4vw, 42px);
    font-weight: 800;
    line-height: 1.2;
    letter-spacing: -1px;
}

.dashboard-hero p {
    position: relative;
    max-width: 680px;
    color: #e3f1ff !important;
    font-size: 16px;
    line-height: 1.8;
}

/* Premium travel overview cards */
div[data-testid="stMetric"] {
    min-height: 125px;
    background: linear-gradient(
        145deg,
        #ffffff,
        #f4f9ff
    );
    border: 1px solid #dceaff;
    border-radius: 20px;
    padding: 22px;
    box-shadow: 0 6px 20px rgba(22, 76, 145, 0.07);
    transition:
        transform 0.25s ease,
        box-shadow 0.25s ease,
        border-color 0.25s ease;
}

div[data-testid="stMetric"]:hover {
    transform: translateY(-5px);
    border-color: #83bbff;
    box-shadow: 0 14px 28px rgba(8, 124, 255, 0.15);
}

div[data-testid="stMetricValue"] {
    color: #087cff;
    font-weight: 800;
}

/* Smooth dashboard entrance */
@keyframes dashboardFadeUp {
    from {
        opacity: 0;
        transform: translateY(12px);
    }
    to {
        opacity: 1;
        transform: translateY(0);
    }
}

/* Respect reduced-motion preferences */
@media (prefers-reduced-motion: reduce) {
    .dashboard-hero {
        animation: none;
    }

    div[data-testid="stMetric"] {
        transition: none;
    }

    div[data-testid="stMetric"]:hover {
        transform: none;
    }
}
    
</style>
""", unsafe_allow_html=True)


# --- Welcome page ---
if st.session_state.page == "Home":

    # Top navigation
    logo, spacer, nav1, nav2 = st.columns([5, 2, 1.2, 1.5])

    with logo:
        st.markdown("## 🧳 TripPlan")

    with spacer:
        st.write("")

    # Keep navigation buttons slightly lower
    with nav1:
        st.write("")
        st.write("")
        st.button(
            "Log In",
            key="nav_login",
            on_click=go_to,
            args=("Login",),
            use_container_width=True
        )

    with nav2:
        st.write("")
        st.write("")
        st.button(
            "Get Started",
            key="nav_signup",
            on_click=go_to,
            args=("Signup",),
            use_container_width=True
        )

    st.divider()

    # Hero section
    left, right = st.columns([1.1, 1], gap="large")

    with left:
        st.write("")
        st.write("")
        st.markdown(
            "<p style='color:#087cff; font-weight:700; "
            "letter-spacing:2px; font-size:13px;'>"
            "YOUR NEXT JOURNEY STARTS HERE</p>",
            unsafe_allow_html=True
        )

        st.markdown(
            "<h1 style='font-size:48px; line-height:1.15; "
            "margin-bottom:18px;'>Explore more.<br>"
            "Worry less.</h1>",
            unsafe_allow_html=True
        )

        st.markdown("""
        Plan your next adventure with confidence.

        Organize your trips, discover incredible destinations,
        and create personalized itineraries — all in one place.
        """)

        st.write("")
        st.write("")

        # Hero buttons
        btn1, btn2 = st.columns(2, gap="small")

        with btn1:
            st.button(
                "✈️ Start Planning",
                key="hero_signup",
                on_click=go_to,
                args=("Signup",),
                use_container_width=True
            )

        with btn2:
            st.button(
                "Log In →",
                key="hero_login",
                on_click=go_to,
                args=("Login",),
                use_container_width=True
            )

        st.write("")
        st.write("")

        # Feature highlights
        feature1, feature2, feature3 = st.columns(3)

        with feature1:
            st.markdown("📍 **Discover**")
            st.caption("Amazing places")

        with feature2:
            st.markdown("🗺️ **Organize**")
            st.caption("Your itinerary")

        with feature3:
            st.markdown("🔔 **Stay ready**")
            st.caption("Travel updates")

    with right:
        st.write("")
        st.write("")

        image_path = os.path.join(
            os.path.dirname(os.path.abspath(__file__)),
            "static",
            "itenarary_planner.jpeg"
        )

        if os.path.isfile(image_path):
            st.image(
                image_path,
                use_container_width=True
            )
        else:
            st.info("Travel planner image will be available soon.")

        st.markdown(
            "<p style='text-align:center; font-size:13px; "
            "color:#7890ad;'>"
            "Your next adventure, beautifully organized."
            "</p>",
            unsafe_allow_html=True
        )

# --- Login page ---
elif st.session_state.page == "Login":

    st.markdown("# Welcome Back! 👋")
    st.write("Log in to access your trips and travel plans")

    with st.form("login_form"):
        email = st.text_input("Email address")
        password = st.text_input("Password", type="password")
        remember = st.checkbox("Remember me")

        submitted = st.form_submit_button("Log In")

        if submitted:
            if not email.strip() or not password:
                st.warning("Please enter both your email and password.")

            elif (
                "@" not in email
                or "." not in email.split("@")[-1]
                or " " in email.strip()
            ):
                st.error("Please enter a valid email address.")

            else:
                st.success("Your login details look valid!")
                st.info(
                    "This is still a prototype. "
                    "Your identity has not been verified because "
                    "real authentication is not connected yet."
                )

    st.button(
        "Create an account",
        on_click=go_to,
        args=("Signup",)
    )

    st.button(
        "← Back to Home",
        on_click=go_to,
        args=("Home",)
    )

    st.button(
        "Open Demo Dashboard",
        key="open_demo_dashboard",
        on_click=go_to,
        args=("Dashboard",)
    )

# --- Sign-up page ---
elif st.session_state.page == "Signup":

    st.markdown("# Join TripPlan 🌍")
    st.write("Create your account and start planning your next adventure.")

    with st.form("signup_form"):
        name = st.text_input("Full name")
        email = st.text_input("Email address")
        password = st.text_input(
            "Create a password",
            type="password"
        )
        confirm = st.text_input(
            "Confirm password",
            type="password"
        )

        submitted = st.form_submit_button("Create Account")

        if submitted:
            # Check for empty fields
            if not name.strip() or not email.strip() or not password or not confirm:
                st.warning("Please fill in all fields.")

            # Check email format
            elif (
                "@" not in email
                or "." not in email.split("@")[-1]
                or " " in email.strip()
            ):
                st.error("Please enter a valid email address.")

            # Check password length
            elif len(password) < 8:
                st.error("Your password must be at least 8 characters long.")

            # Check that passwords match
            elif password != confirm:
                st.error("Your passwords don't match.")

            # Validation passed
            else:
                st.success("All fields passed validation!")
                st.info(
                    "This is still a prototype. "
                    "Your account has not been created because "
                    "registration and secure password storage "
                    "are not connected yet."
                )

    st.button(
        "Already have an account? Log in",
        on_click=go_to,
        args=("Login",)
    )
    st.button(
        "← Back to Home",
        on_click=go_to,
        args=("Home",)
    )

# --- Dashboard page ---
elif st.session_state.page == "Dashboard":

    if "trips" not in st.session_state:
        st.session_state.trips = []

    if "wishlist" not in st.session_state:
        st.session_state.wishlist = []



    # Dashboard header navigation
    logo, spacer, home_col, logout_col = st.columns(
        [4, 2, 1.2, 1.3],
        vertical_alignment="center"
    )

    with logo:
        st.markdown("## 🧳 TripPlan")

    with spacer:
        st.write("")

    with home_col:
        st.button(
            "🏠 Home",
            key="dashboard_home",
            on_click=go_to,
            args=("Home",),
            use_container_width=True
        )

    with logout_col:
        st.button(
            "↩️ Log Out",
            key="dashboard_logout",
            on_click=go_to,
            args=("Home",),
            use_container_width=True
        )

    st.divider()

    # Welcome banner
    st.markdown(
        """
        <div class="dashboard-hero">
            <span class="hero-label">
                ✦ YOUR NEXT ADVENTURE STARTS HERE
            </span>
            <h1>Make every journey unforgettable. 🌍</h1>
            <p>Your next adventure starts with a great plan.
            Discover destinations, organize your itineraries,
            and keep all your travel inspiration in one place.</p>
        </div>
        """,
        unsafe_allow_html=True
    )

    # Overview cards

    # Overview cards
    st.markdown(
        '<div class="dashboard-section">Your Travel Overview</div>',
        unsafe_allow_html=True
    )

    col1, col2, col3 = st.columns(3, gap="medium")

    with col1:
        st.metric(
            label="✈️ My Trips",
            value=len(st.session_state.trips)
        )

    with col2:
        st.metric(
            label="🌍 Dream Destinations",
            value=len(st.session_state.wishlist)
        )

    with col3:
        st.metric(
            label="🗺️ Travel Planner",
            value="Ready"
        )


    # --- Destination Discovery ---
    st.markdown(
        '<div class="dashboard-section">'
        '✨ Find Your Next Adventure</div>',
        unsafe_allow_html=True
    )

    st.write(
        "Need inspiration? Explore these destinations "
        "and save your favourites to your wishlist."
    )

    inspiration_destinations = [
        {
            "name": "Kyoto, Japan",
            "emoji": "🏯",
            "description": "Temples, gardens and beautiful traditions.",
            "key": "kyoto",
            "image": "https://images.unsplash.com/photo-1493976040374-85c8e12f0c0e?auto=format&fit=crop&w=900&h=600&q=85",
            "details": "Discover peaceful Japanese gardens, historic temples, traditional tea houses and the famous Arashiyama Bamboo Grove."
        },
        {
            "name": "Bali, Indonesia",
            "emoji": "🌴",
            "description": "Tropical beaches and peaceful escapes.",
            "key": "bali",
            "image": "https://images.unsplash.com/photo-1537996194471-e657df975ab4?auto=format&fit=crop&w=900&h=600&q=85",
            "details": "Explore tropical beaches, lush rice terraces, beautiful waterfalls and the cultural attractions around Ubud."
        },
        {
            "name": "Paris, France",
            "emoji": "🗼",
            "description": "Iconic landmarks, art and cafés.",
            "key": "paris",
            "image": "https://images.unsplash.com/photo-1502602898657-3e91760cbb34?auto=format&fit=crop&w=900&h=600&q=85",
            "details": "Visit the Eiffel Tower, explore world-famous museums, stroll along the Seine and enjoy the city's charming cafés."
        },
        {
            "name": "Cape Town, South Africa",
            "emoji": "🌊",
            "description": "Mountain views and coastal adventures.",
            "key": "cape_town",
            "image": "https://images.unsplash.com/photo-1580060839134-75a5edca2e99?auto=format&fit=crop&w=900&h=600&q=85",
            "details": "Enjoy views of Table Mountain, explore the V&A Waterfront, visit the Cape Peninsula and discover beautiful coastal scenery."
        },
        {
            "name": "Rome, Italy",
            "emoji": "🏛️",
            "description": "Ancient history and amazing architecture.",
            "key": "rome",
            "image": "https://images.unsplash.com/photo-1552832230-c0197dd311b5?auto=format&fit=crop&w=900&h=600&q=85",
            "details": "Step into history at the Colosseum and Roman Forum, admire historic piazzas and discover the city's famous Italian cuisine."
        },
        {
            "name": "Santorini, Greece",
            "emoji": "🌅",
            "description": "Blue seas and stunning island views.",
            "key": "santorini",
            "image": "https://images.unsplash.com/photo-1613395877344-13d4a8e0d49e?auto=format&fit=crop&w=900&h=600&q=85",
            "details": "Explore whitewashed villages, admire the blue-domed churches, relax by the Aegean Sea and enjoy spectacular sunset views."
        }
    ]

    for start in range(0, len(inspiration_destinations), 3):
        destination_row = inspiration_destinations[start:start + 3]
        columns = st.columns(3, gap="medium")

        for column, destination in zip(columns, destination_row):
            with column:
                with st.container(border=True):
                    st.image(
                        destination["image"],
                        use_container_width=True
                    )

                    st.markdown(
                        f"### {destination['emoji']} "
                        f"{destination['name']}"
                    )

                    st.caption(destination["description"])

                    if st.button(
                            "Explore Destination",
                            key=f"explore_{destination['key']}",
                            use_container_width=True
                    ):
                        st.session_state["selected_destination"] = destination["key"]

                    if st.button(
                        "♡ Save to Wishlist",
                        key=f"save_inspiration_{destination['key']}",
                        use_container_width=True
                    ):
                        destination_name = destination["name"]

                        already_saved = any(
                            (
                                item.get("name", "").casefold()
                                if isinstance(item, dict)
                                else str(item).casefold()
                            ) == destination_name.casefold()
                            for item in st.session_state.wishlist
                        )

                        if already_saved:
                            st.info("This destination is already saved!")
                        else:
                            st.session_state.wishlist.append(
                                destination_name
                            )
                            st.success("Destination saved! 💙")

    st.write("")
    st.divider()

    # --- End Destination Discovery ---

    # --- Selected Destination Details ---
    selected_key = st.session_state.get("selected_destination")

    if selected_key:
        selected_destination = next(
            (
                destination
                for destination in inspiration_destinations
                if destination["key"] == selected_key
            ),
            None
        )

        if selected_destination:
            st.divider()
            st.subheader(
                f"{selected_destination['emoji']} "
                f"{selected_destination['name']}"
            )

            st.image(
                selected_destination["image"],
                use_container_width=True
            )

            st.write(selected_destination["details"])

            if st.button(
                    "✈️ Plan a Trip to This Destination",
                    key="plan_selected_destination",
                    use_container_width=True
            ):
                st.session_state["new_trip_destination"] = (
                    selected_destination["name"]
                )
                st.session_state["show_trip_form"] = True
                st.rerun()

            if st.button("Close Details", key="close_destination_details"):
                del st.session_state["selected_destination"]
                st.rerun()

    # --- End Selected Destination Details ---

    # --- Interactive Travel Map ---

    st.markdown("## 🗺️ Your Travel Map")
    st.write(
        "Explore your planned trips and wishlist destinations "
        "on one interactive map."
    )

    # Collect destinations from saved trips and wishlist
    map_locations = []

    for trip in st.session_state.trips:
        destination = trip.get("destination", "").strip()

        if destination:
            map_locations.append({
                "destination": destination,
                "category": "Planned Trip",
                "details": trip.get("dates", ""),
                "color": "blue"
            })

    for place in st.session_state.wishlist:
        if isinstance(place, dict):
            map_locations.append({
                "destination": place.get("name", "Saved place"),
                "category": "Wishlist",
                "details": "Dream destination",
                "color": "green",
                "latitude": place.get("latitude"),
                "longitude": place.get("longitude"),
                "country": place.get("country", ""),
            })
        elif isinstance(place, str) and place.strip():
            # Support wishlist entries saved before real-place search was added.
            map_locations.append({
                "destination": place.strip(),
                "category": "Wishlist",
                "details": "Dream destination",
                "color": "green"
            })

    # Let users choose which pins to display
    map_filter = st.radio(
        "Show on map",
        ["All Places", "Planned Trips", "Wishlist"],
        horizontal=True,
        key="travel_map_filter"
    )

    if map_filter == "Planned Trips":
        visible_locations = [
            place for place in map_locations
            if place["category"] == "Planned Trip"
        ]
    elif map_filter == "Wishlist":
        visible_locations = [
            place for place in map_locations
            if place["category"] == "Wishlist"
        ]
    else:
        visible_locations = map_locations

    if not visible_locations:
        st.info(
            "Add a trip or a wishlist destination to see it on your map."
        )
    else:
        coordinates_found = []
        failed_destinations = []

        with st.spinner("Finding your destinations..."):
            for place in visible_locations:
                try:
                    if (
                        place.get("latitude") is not None
                        and place.get("longitude") is not None
                    ):
                        coords = {
                            "latitude": place["latitude"],
                            "longitude": place["longitude"],
                            "name": place["destination"],
                            "country": place.get("country", ""),
                        }
                    else:
                        coords = get_destination_coordinates(
                            place["destination"]
                        )

                    if coords:
                        coordinates_found.append({
                            **place,
                            **coords
                        })
                    else:
                        failed_destinations.append(
                            place["destination"]
                        )

                except requests.RequestException:
                    failed_destinations.append(
                        place["destination"]
                    )

        if coordinates_found:
            # Center the map around the saved destinations
            average_lat = sum(
                place["latitude"] for place in coordinates_found
            ) / len(coordinates_found)

            average_lon = sum(
                place["longitude"] for place in coordinates_found
            ) / len(coordinates_found)

            travel_map = folium.Map(
                location=[average_lat, average_lon],
                zoom_start=3,
                tiles="OpenStreetMap"
            )

            for place in coordinates_found:
                popup_html = f"""
                <div style="width: 200px;">
                    <b>{place['name']}, {place['country']}</b><br>
                    <b>Type:</b> {place['category']}<br>
                    <b>Details:</b> {place['details']}
                </div>
                """

                folium.Marker(
                    location=[
                        place["latitude"],
                        place["longitude"]
                    ],
                    tooltip=place["name"],
                    popup=folium.Popup(
                        popup_html,
                        max_width=250
                    ),
                    icon=folium.Icon(
                        color=place["color"],
                        icon="map-marker",
                        prefix="fa"
                    )
                ).add_to(travel_map)

            st_folium(
                travel_map,
                use_container_width=True,
                height=500,
                key="tripplan_interactive_map"
            )

            st.caption(
                "🔵 Blue pins = planned trips  |  "
                "🟢 Green pins = wishlist destinations"
            )

            st.write(
                f"**{len(coordinates_found)} location(s) "
                "shown on your map.**"
            )

        else:
            st.warning(
                "We couldn't locate these destinations. "
                "Try using a city name and country."
            )

        if failed_destinations:
            st.warning(
                "Could not locate: "
                + ", ".join(failed_destinations)
                + ". Check the spelling or add the country name."
            )

    st.markdown("## 🌦️ Destination Weather")
    st.write("Check current conditions and the forecast for your next adventure.")

    if st.session_state.trips:
        weather_destinations = list(dict.fromkeys(
            trip["destination"]
            for trip in st.session_state.trips
            if trip.get("destination", "").strip()
        ))
    else:
        weather_destinations = []

    weather_options = weather_destinations + ["Other destination"]

    selected_weather_destination = st.selectbox(
        "Choose a destination",
        weather_options,
        key="weather_destination"
    )

    if selected_weather_destination == "Other destination":
        weather_city = st.text_input(
            "Enter a city",
            placeholder="e.g. Cape Town, South Africa",
            key="weather_city"
        )
    else:
        weather_city = selected_weather_destination

    if st.button("🌤️ Get Weather", key="get_weather_button"):
        if not weather_city.strip():
            st.warning("Please enter a destination.")
        else:
            try:
                with st.spinner("Getting the latest forecast..."):
                    result = get_weather(weather_city.strip())

                if result is None:
                    st.warning(
                        "We couldn't find that destination. "
                        "Try entering a city and country."
                    )
                else:
                    place = result["place"]
                    weather = result["weather"]
                    current = weather["current"]
                    daily = weather["daily"]

                    st.markdown(
                        f"### 📍 {place['name']}, "
                        f"{place.get('country', '')}"
                    )

                    st.caption("Current conditions")

                    st.markdown(
                        f"### {current['temperature_2m']}°C"
                    )
                    st.write(
                        weather_description(current["weather_code"])
                    )

                    col1, col2, col3 = st.columns(3)

                    col1.metric(
                        "Feels like",
                        f"{current['apparent_temperature']}°C"
                    )
                    col2.metric(
                        "Humidity",
                        f"{current['relative_humidity_2m']}%"
                    )
                    col3.metric(
                        "Wind",
                        f"{current['wind_speed_10m']} km/h"
                    )

                    st.markdown("### 🗓️ 7-Day Forecast")

                    forecast = pd.DataFrame({
                        "Date": daily["time"],
                        "Conditions": [
                            weather_description(code)
                            for code in daily["weather_code"]
                        ],
                        "High (°C)": daily["temperature_2m_max"],
                        "Low (°C)": daily["temperature_2m_min"],
                        "Rain chance (%)": [
                            value if value is not None else "N/A"
                            for value in daily[
                                "precipitation_probability_max"
                            ]
                        ]
                    })

                    st.dataframe(
                        forecast,
                        use_container_width=True,
                        hide_index=True
                    )

                    st.caption(
                        "Weather data provided by Open-Meteo. "
                        "Forecasts can change."
                    )

            except requests.RequestException:
                st.error(
                    "Weather data could not be loaded. "
                    "Check your internet connection and try again."
                )

    st.write("")
    st.divider()

    # Create a trip
    st.markdown(
        '<div class="dashboard-section">✈️ Plan a New Trip</div>',
        unsafe_allow_html=True
    )

    with st.container(border=True):
        st.write("Where will your next adventure take you?")

        with st.form("new_trip_form"):

            destination = st.text_input(
                "Destination",
                placeholder="e.g. Cape Town, Paris, Tokyo",
                key="new_trip_destination"
            )

            trip_dates = st.text_input(
                "Travel dates",
                placeholder="e.g. 10–17 December 2026"
            )

            trip_notes = st.text_area(
                "Trip notes",
                placeholder="Add activities, ideas, or things to remember..."
            )

            submitted = st.form_submit_button(
                "＋ Create My Trip",
                use_container_width=True
            )


            if submitted:
                clean_destination = destination.strip()
                clean_dates = trip_dates.strip()
                clean_notes = trip_notes.strip()

                if not clean_destination:
                    st.warning("Please enter a destination.")

                elif not clean_dates:
                    st.warning("Please enter your travel dates.")

                elif any(
                    trip.get("destination", "").casefold() == clean_destination.casefold()
                    and trip.get("dates", "").strip().casefold() == clean_dates.casefold()
                    for trip in st.session_state.trips
                ):
                    st.warning(
                        "You already have a trip to this destination "
                        "with these dates."
                    )

                else:

                    st.session_state.trips.append({
                        "destination": clean_destination,
                        "dates": clean_dates,
                        "notes": clean_notes,
                        "itinerary": [],
                        "budget": 0.0,
                        "currency": "ZAR",
                        "expenses": []
                    })

                    st.success(
                        f"Your trip to {clean_destination} was created!"
                    )
                    save_trip_data()
                    st.rerun()

    st.write("")


    # --- TripPlan Analytics Dashboard ---
    st.markdown("## 📈 Your Travel Analytics")
    st.caption("A live overview of your trips, budgets and spending.")

    trips_for_analytics = st.session_state.get("trips", [])

    analytics_rows = []
    category_rows = []

    for trip in trips_for_analytics:
        destination_name = str(trip.get("destination", "Unknown"))
        trip_currency = str(trip.get("currency", "ZAR"))

        try:
            trip_budget = float(trip.get("budget", 0) or 0)
        except (ValueError, TypeError):
            trip_budget = 0.0

        trip_spent = 0.0

        for expense in trip.get("expenses", []) or []:
            try:
                amount = float(expense.get("amount", 0) or 0)
            except (ValueError, TypeError):
                amount = 0.0

            expense_currency = str(
                expense.get("currency", trip_currency)
            )

            category = str(expense.get("category", "Other"))

            category_rows.append({
                "Destination": destination_name,
                "Currency": expense_currency,
                "Category": category,
                "Amount": amount,
            })

            # Only compare expenses in the trip's own currency.
            if expense_currency == trip_currency:
                trip_spent += amount

        analytics_rows.append({
            "Destination": destination_name,
            "Currency": trip_currency,
            "Budget": trip_budget,
            "Spent": trip_spent,
            "Remaining": trip_budget - trip_spent,
        })

    total_expense_count = sum(
        len(trip.get("expenses", []) or [])
        for trip in trips_for_analytics
    )

    metric1, metric2, metric3 = st.columns(3)

    metric1.metric("✈️ Total Trips", len(trips_for_analytics))
    metric2.metric("🧾 Expense Entries", total_expense_count)
    metric3.metric(
        "📍 Wishlist Places",
        len(st.session_state.get("wishlist", []))
    )

    if not analytics_rows:
        st.info(
            "🌍 Your analytics will appear here once you create "
            "your first trip. Start planning your next adventure!"
        )
    else:
        analytics_df = pd.DataFrame(analytics_rows)

        st.markdown("### 💰 Budget vs Spending")

        available_currencies = sorted(
            analytics_df["Currency"].dropna().unique().tolist()
        )

        if len(available_currencies) > 1:
            st.caption(
                "Choose a currency to compare trips fairly. "
                "Different currencies are never combined."
            )

        selected_currency = st.selectbox(
            "Analytics currency",
            options=available_currencies,
            key="analytics_currency"
        )

        currency_df = analytics_df[
            analytics_df["Currency"] == selected_currency
        ].copy()

        chart_df = currency_df.melt(
            id_vars=["Destination"],
            value_vars=["Budget", "Spent", "Remaining"],
            var_name="Measure",
            value_name="Amount"
        )

        fig_budget = px.bar(
            chart_df,
            x="Destination",
            y="Amount",
            color="Measure",
            barmode="group",
            title=f"Your {selected_currency} Budget Overview",
            color_discrete_map={
                "Budget": "#17365D",
                "Spent": "#18B6A4",
                "Remaining": "#8BD9D0",
            },
            template="plotly_white",
        )

        fig_budget.update_layout(
            legend_title_text="",
            margin=dict(l=20, r=20, t=60, b=20),
            font=dict(family="Arial", color="#10243A"),
        )

        st.plotly_chart(fig_budget, use_container_width=True)

        st.markdown("### 🧾 Spending by Category")

        if category_rows:
            category_df = pd.DataFrame(category_rows)
            category_currency_df = category_df[
                category_df["Currency"] == selected_currency
            ]

            if not category_currency_df.empty:
                category_summary = (
                    category_currency_df.groupby(
                        "Category", as_index=False
                    )["Amount"].sum()
                )

                fig_category = px.pie(
                    category_summary,
                    names="Category",
                    values="Amount",
                    hole=0.58,
                    title=f"Spending Breakdown ({selected_currency})",
                    color_discrete_sequence=[
                        "#18B6A4", "#17365D", "#45C4D4",
                        "#85A9C9", "#A8E6DE", "#64748B"
                    ],
                )

                fig_category.update_layout(
                    margin=dict(l=20, r=20, t=60, b=20),
                    font=dict(family="Arial", color="#10243A"),
                )

                st.plotly_chart(
                    fig_category,
                    use_container_width=True
                )
            else:
                st.info(
                    f"No expense categories recorded in {selected_currency} yet."
                )
        else:
            st.info(
                "🧳 No expenses yet. Add an expense to see your "
                "spending breakdown here."
            )

        st.markdown("### 📋 Trip Budget Summary")

        summary_df = currency_df[
            ["Destination", "Currency", "Budget", "Spent", "Remaining"]
        ].copy()

        st.dataframe(
            summary_df,
            use_container_width=True,
            hide_index=True,
        )

    # --- End TripPlan Analytics Dashboard ---
    # --- Excel Export Centre ---
    st.markdown("## 📊 Your TripPlan Excel Centre")

    st.write(
        "Download your trips, budgets, expenses, itinerary, "
        "and wishlist in one organised Excel workbook."
    )

    excel_file = create_tripplan_excel()

    st.download_button(
        label="📥 Download All My TripPlan Data",
        data=excel_file,
        file_name="TripPlan_Travel_Report.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        use_container_width=True,
        key="download_tripplan_excel"
    )
    # --- End Excel Export Centre ---


    # My Trips and Day-by-Day Itinerary
    st.markdown(
        '<div class="dashboard-section">🗺️ My Trips & Itineraries</div>',
        unsafe_allow_html=True
    )

    if st.session_state.trips:

        for i, trip in enumerate(st.session_state.trips):

            # Support trips created before the itinerary feature
            if "itinerary" not in trip:
                trip["itinerary"] = []

            with st.container(border=True):

                st.markdown(f"### 📍 {trip['destination']}")

                if trip["dates"]:
                    st.caption(f"📅 Trip dates: {trip['dates']}")

                if trip["notes"]:
                    st.write(trip["notes"])

                st.divider()


                # --- Travel Budget Planner ---
                st.markdown("#### 💰 Travel Budget")

                # Support trips created before the budget feature
                if "budget" not in trip:
                    trip["budget"] = 0.0
                if "currency" not in trip:
                    trip["currency"] = "ZAR"
                if "expenses" not in trip:
                    trip["expenses"] = []

                currency_symbols = {
                    "ZAR": "R",
                    "USD": "$",
                    "EUR": "€",
                    "GBP": "£",
                    "JPY": "¥"
                }

                currency_symbol = currency_symbols.get(
                    trip["currency"], trip["currency"] + " "
                )


                budget_col, currency_col = st.columns(2)

                with budget_col:
                    st.metric(
                        "Total trip budget",
                        f"{currency_symbols.get(trip['currency'], trip['currency'] + ' ')}"
                        f"{float(trip['budget']):,.2f}"
                    )

                with currency_col:
                    currency_options = [
                        "ZAR", "USD", "EUR", "GBP", "JPY"
                    ]

                    selected_currency = st.selectbox(
                        "Currency",
                        currency_options,
                        index=currency_options.index(trip["currency"])
                        if trip["currency"] in currency_options else 0,
                        key=f"trip_currency_{i}"
                    )

                trip["currency"] = selected_currency
                currency_symbol = currency_symbols[selected_currency]

                with st.expander("✏️ Edit Trip Budget"):
                    with st.form(f"edit_budget_form_{i}"):
                        edited_budget = st.number_input(
                            "New total trip budget",
                            min_value=0.0,
                            value=float(trip["budget"]),
                            step=100.0,
                            key=f"edit_budget_value_{i}"
                        )

                        budget_saved = st.form_submit_button(
                            "💾 Save Budget",
                            use_container_width=True
                        )

                        if budget_saved:
                            trip["budget"] = float(edited_budget)
                            st.success("Trip budget updated!")
                            st.rerun()

                new_budget = float(trip["budget"])


                total_spent = sum(
                    float(expense.get("amount", 0))
                    for expense in trip["expenses"]
                )
                remaining_budget = new_budget - total_spent

                metric1, metric2, metric3 = st.columns(3)

                metric1.metric(
                    "Budget",
                    f"{currency_symbol}{new_budget:,.2f}"
                )
                metric2.metric(
                    "Spent",
                    f"{currency_symbol}{total_spent:,.2f}"
                )
                metric3.metric(
                    "Remaining",
                    f"{currency_symbol}{remaining_budget:,.2f}",
                    delta_color="normal"
                )


                if new_budget > 0:
                    progress = total_spent / new_budget
                    st.progress(min(progress, 1.0))

                    if total_spent > new_budget:
                        st.error(
                            f"🚨 Over budget by "
                            f"{currency_symbol}"
                            f"{total_spent - new_budget:,.2f}"
                        )
                    elif progress >= 0.8:
                        st.warning(
                            "⚠️ Budget alert: you've used "
                            f"{progress:.0%} of your budget."
                        )
                    else:
                        st.success(
                            f"✅ You're on track! "
                            f"You've used {progress:.0%} "
                            "of your budget."
                        )
                else:
                    st.info(
                        "Set a trip budget above zero to track "
                        "your spending progress."
                    )


                st.markdown("##### ➕ Add an Expense")

                with st.form(f"expense_form_{i}"):
                    expense_name = st.text_input(
                        "Expense description",
                        placeholder="e.g. Hotel booking",
                        key=f"expense_name_{i}"
                    )

                    expense_col1, expense_col2 = st.columns(2)

                    with expense_col1:
                        expense_amount = st.number_input(
                            "Amount",
                            min_value=0.0,
                            step=10.0,
                            key=f"expense_amount_{i}"
                        )

                    with expense_col2:
                        expense_category = st.selectbox(
                            "Category",
                            [
                                "Flights",
                                "Accommodation",
                                "Food",
                                "Transport",
                                "Activities",
                                "Shopping",
                                "Other"
                            ],
                            key=f"expense_category_{i}"
                        )

                    expense_submitted = st.form_submit_button(
                        "＋ Save Expense",
                        use_container_width=True
                    )

                    if expense_submitted:
                        if not expense_name.strip():
                            st.warning("Please enter an expense description.")
                        elif expense_amount <= 0:
                            st.warning("Please enter an amount greater than zero.")
                        else:
                            trip["expenses"].append({
                                "name": expense_name.strip(),
                                "amount": float(expense_amount),
                                "category": expense_category
                            })
                            st.success("Expense saved!")
                            st.rerun()


                if trip["expenses"]:
                    st.markdown("##### 🧾 Your Expenses")

                    expense_rows = [
                        {
                            "Expense": expense["name"],
                            "Category": expense["category"],
                            "Amount": (
                                f"{currency_symbol}"
                                f"{float(expense['amount']):,.2f}"
                            )
                        }
                        for expense in trip["expenses"]
                    ]

                    st.dataframe(
                        pd.DataFrame(expense_rows),
                        use_container_width=True,
                        hide_index=True
                    )

                    # Spending breakdown by category
                    st.markdown("##### 📊 Spending by Category")

                    category_totals = {}

                    for expense in trip["expenses"]:
                        category = expense["category"]
                        amount = float(expense["amount"])

                        category_totals[category] = (
                            category_totals.get(category, 0) + amount
                        )

                    category_df = pd.DataFrame({
                        "Category": list(category_totals.keys()),
                        "Amount": list(category_totals.values())
                    })

                    st.bar_chart(
                        category_df,
                        x="Category",
                        y="Amount"
                    )

                    for category, amount in category_totals.items():
                        st.write(
                            f"**{category}:** "
                            f"{currency_symbol}{amount:,.2f}"
                        )



                    st.markdown("##### ✏️ Edit or Delete Expenses")

                    expense_categories = [
                        "Flights",
                        "Accommodation",
                        "Food",
                        "Transport",
                        "Activities",
                        "Shopping",
                        "Other"
                    ]

                    for expense_index, expense in enumerate(
                        trip["expenses"]
                    ):
                        with st.expander(
                            f"🧾 {expense['name']} — "
                            f"{currency_symbol}"
                            f"{float(expense['amount']):,.2f}"
                        ):
                            with st.form(
                                f"edit_expense_form_{i}_{expense_index}"
                            ):
                                edited_name = st.text_input(
                                    "Expense description",
                                    value=expense["name"],
                                    key=f"edit_name_{i}_{expense_index}"
                                )

                                edited_amount = st.number_input(
                                    "Amount",
                                    min_value=0.01,
                                    value=max(
                                        0.01,
                                        float(expense["amount"])
                                    ),
                                    step=10.0,
                                    key=f"edit_amount_{i}_{expense_index}"
                                )

                                current_category = expense.get(
                                    "category", "Other"
                                )

                                edited_category = st.selectbox(
                                    "Category",
                                    expense_categories,
                                    index=(
                                        expense_categories.index(
                                            current_category
                                        )
                                        if current_category in expense_categories
                                        else expense_categories.index("Other")
                                    ),
                                    key=f"edit_category_{i}_{expense_index}"
                                )

                                edit_col, delete_col = st.columns(2)

                                with edit_col:
                                    save_expense = st.form_submit_button(
                                        "💾 Save Changes",
                                        use_container_width=True
                                    )

                                with delete_col:
                                    delete_expense = st.form_submit_button(
                                        "🗑️ Delete Expense",
                                        use_container_width=True
                                    )

                                if save_expense:
                                    if not edited_name.strip():
                                        st.warning(
                                            "Please enter an expense description."
                                        )
                                    else:
                                        expense["name"] = edited_name.strip()
                                        expense["amount"] = float(
                                            edited_amount
                                        )
                                        expense["category"] = edited_category

                                        st.success("Expense updated!")
                                        st.rerun()

                                if delete_expense:
                                    trip["expenses"].pop(expense_index)
                                    st.rerun()



                else:
                    st.caption("No expenses added yet.")


                st.divider()
                # --- End Travel Budget Planner ---


                # Add an activity
                st.markdown("#### ➕ Add an Activity")

                with st.form(f"activity_form_{i}"):

                    activity_name = st.text_input(
                        "Activity name",
                        placeholder="e.g. Visit the Eiffel Tower",
                        key=f"activity_name_{i}"
                    )

                    date_col, time_col = st.columns(2)

                    with date_col:
                        activity_date = st.date_input(
                            "Activity date",
                            key=f"activity_date_{i}"
                        )

                    with time_col:
                        activity_time = st.time_input(
                            "Start time",
                            value=time(9, 0),
                            key=f"activity_time_{i}"
                        )

                    activity_notes = st.text_area(
                        "Activity notes",
                        placeholder="Address, booking details, or things to remember...",
                        key=f"activity_notes_{i}"
                    )

                    add_activity = st.form_submit_button(
                        "＋ Add to Itinerary",
                        use_container_width=True
                    )


                    if add_activity:
                        clean_activity_name = activity_name.strip()
                        activity_date_str = activity_date.isoformat()
                        activity_time_str = activity_time.strftime("%H:%M")

                        if not clean_activity_name:
                            st.warning("Please enter an activity name.")

                        elif any(
                            activity.get("name", "").casefold()
                            == clean_activity_name.casefold()
                            and activity.get("date") == activity_date_str
                            and activity.get("time") == activity_time_str
                            for activity in trip["itinerary"]
                        ):
                            st.warning(
                                "This activity already exists at that "
                                "date and time."
                            )

                        else:
                            trip["itinerary"].append({
                                "name": clean_activity_name,
                                "date": activity_date_str,
                                "time": activity_time_str,
                                "notes": activity_notes.strip()
                            })

                            trip["itinerary"].sort(
                                key=lambda activity: (
                                    activity["date"],
                                    activity["time"]
                                )
                            )

                            st.success("Activity added to your itinerary!")
                            st.rerun()



                # Day tabs, weekly calendar, and visual timeline
                st.markdown("#### 📅 Your Itinerary")

                if trip["itinerary"]:

                    # Sort activities by date and time
                    trip["itinerary"].sort(
                        key=lambda activity: (
                            activity["date"],
                            activity["time"]
                        )
                    )

                    itinerary_dates = sorted(set(
                        activity["date"]
                        for activity in trip["itinerary"]
                    ))

                    # Three ways to view the same itinerary
                    day_view, week_view, timeline_view = st.tabs([
                        "📆 Day View",
                        "🗓️ Week View",
                        "🕒 Timeline"
                    ])

                    # ----------------------------------
                    # 1. DAY VIEW — preserve existing tabs
                    # ----------------------------------
                    with day_view:

                        day_tabs = st.tabs([
                            date.fromisoformat(day).strftime("%a, %d %b")
                            for day in itinerary_dates
                        ])

                        for day, day_tab in zip(
                            itinerary_dates, day_tabs
                        ):

                            with day_tab:
                                st.caption(
                                    date.fromisoformat(day).strftime(
                                        "%A, %d %B %Y"
                                    )
                                )

                                day_activities = [
                                    (j, activity)
                                    for j, activity in enumerate(
                                        trip["itinerary"]
                                    )
                                    if activity["date"] == day
                                ]

                                for j, activity in day_activities:

                                    with st.container(border=True):
                                        time_col, details_col, delete_col = (
                                            st.columns([1, 4, 0.7])
                                        )

                                        with time_col:
                                            st.markdown(
                                                f"**{activity['time']}**"
                                            )

                                        with details_col:
                                            st.markdown(
                                                f"**{activity['name']}**"
                                            )

                                            if activity["notes"]:
                                                st.caption(
                                                    activity["notes"]
                                                )

                                        with delete_col:
                                            if st.button(
                                                "✕",
                                                key=f"delete_activity_{i}_{j}",
                                                help="Delete this activity"
                                            ):
                                                trip["itinerary"].pop(j)
                                                st.rerun()

                    # ----------------------------------
                    # 2. WEEK VIEW — Monday to Sunday
                    # ----------------------------------
                    with week_view:

                        st.write(
                            "Choose a date to see its entire week."
                        )

                        week_anchor = st.date_input(
                            "Choose a date in the week",
                            value=date.today(),
                            key=f"week_picker_{i}"
                        )

                        week_start = week_anchor - timedelta(
                            days=week_anchor.weekday()
                        )

                        week_end = week_start + timedelta(days=6)

                        st.markdown(
                            f"### {week_start.strftime('%d %b')} "
                            f"— {week_end.strftime('%d %b %Y')}"
                        )

                        for day_offset in range(7):

                            current_day = week_start + timedelta(
                                days=day_offset
                            )

                            current_day_str = current_day.isoformat()

                            day_activities = [
                                activity
                                for activity in trip["itinerary"]
                                if activity["date"] == current_day_str
                            ]

                            with st.container(border=True):

                                st.markdown(
                                    f"**{current_day.strftime('%A, %d %b')}**"
                                )

                                if day_activities:

                                    for activity in day_activities:
                                        st.markdown(
                                            f"🕒 **{activity['time']}** "
                                            f"— {activity['name']}"
                                        )

                                        if activity["notes"]:
                                            st.caption(activity["notes"])

                                else:
                                    st.caption("No activities planned.")

                    # ----------------------------------
                    # 3. TIMELINE — visual daily schedule
                    # ----------------------------------
                    with timeline_view:

                        selected_day = st.selectbox(
                            "Choose a day",
                            options=itinerary_dates,
                            format_func=lambda day: (
                                date.fromisoformat(day).strftime(
                                    "%A, %d %B %Y"
                                )
                            ),
                            key=f"timeline_day_{i}"
                        )

                        timeline_activities = [
                            (j, activity)
                            for j, activity in enumerate(
                                trip["itinerary"]
                            )
                            if activity["date"] == selected_day
                        ]

                        st.write("")

                        for position, (j, activity) in enumerate(
                            timeline_activities
                        ):

                            time_col, line_col, details_col = st.columns(
                                [1.1, 0.35, 5]
                            )

                            with time_col:
                                st.markdown(
                                    f"**{activity['time']}**"
                                )

                            with line_col:
                                st.markdown(
                                    "<div style='text-align:center;"
                                    "color:#087cff;font-size:24px;"
                                    "line-height:1.1;'>●</div>",
                                    unsafe_allow_html=True
                                )

                                if position < len(timeline_activities) - 1:
                                    st.markdown(
                                        "<div style='height:55px;"
                                        "width:2px;background:#b9d9ff;"
                                        "margin:auto;'></div>",
                                        unsafe_allow_html=True
                                    )

                            with details_col:
                                with st.container(border=True):
                                    st.markdown(
                                        f"**{activity['name']}**"
                                    )

                                    if activity["notes"]:
                                        st.caption(activity["notes"])

                                    if st.button(
                                        "Delete activity",
                                        key=f"delete_timeline_{i}_{j}"
                                    ):
                                        trip["itinerary"].pop(j)
                                        st.rerun()

                else:
                    st.info(
                        "Your itinerary is empty. "
                        "Add your first activity using the form above!"
                    )


                # Edit Trip button
                if st.button("✏️ Edit Trip", key=f"edit_trip_{i}"):
                    st.session_state[f"editing_trip_{i}"] = True

                # Show the edit form independently of Delete Trip
                if st.session_state.get(f"editing_trip_{i}", False):
                    with st.form(f"edit_trip_form_{i}"):
                        st.subheader("✏️ Edit Your Trip")

                        edited_destination = st.text_input(
                            "Destination",
                            value=trip["destination"],
                            key=f"edit_destination_{i}"
                        )

                        edited_dates = st.text_input(
                            "Trip dates",
                            value=trip.get("dates", ""),
                            key=f"edit_dates_{i}"
                        )

                        edited_notes = st.text_area(
                            "Trip notes",
                            value=trip.get("notes", ""),
                            key=f"edit_notes_{i}"
                        )

                        save_changes = st.form_submit_button("💾 Save Changes")
                        cancel_edit = st.form_submit_button("Cancel")

                    if save_changes:
                        if not edited_destination.strip():
                            st.warning("Please enter a destination.")
                        else:
                            trip["destination"] = edited_destination.strip()
                            trip["dates"] = edited_dates.strip()
                            trip["notes"] = edited_notes.strip()

                            st.session_state[f"editing_trip_{i}"] = False
                            st.success("Trip updated!")
                            st.rerun()

                    if cancel_edit:
                        st.session_state[f"editing_trip_{i}"] = False
                        st.rerun()

                # Delete Trip button — separate from editing
                if st.button("🗑️ Delete Trip", key=f"delete_trip_{i}"):
                    st.session_state[f"confirm_delete_trip_{i}"] = True

                # Ask for confirmation before deleting
                if st.session_state.get(f"confirm_delete_trip_{i}", False):
                    st.warning(
                        f"Are you sure you want to delete "
                        f"your trip to {trip['destination']}?"
                    )

                    confirm_col, cancel_col = st.columns(2)

                    with confirm_col:
                        if st.button(
                            "Yes, delete trip",
                            key=f"confirm_delete_trip_button_{i}"
                        ):
                            st.session_state.trips.pop(i)
                            st.session_state.pop(
                                f"confirm_delete_trip_{i}", None
                            )
                            st.session_state.pop(
                                f"editing_trip_{i}", None
                            )
                            st.rerun()

                    with cancel_col:
                        if st.button(
                            "Cancel deletion",
                            key=f"cancel_delete_trip_{i}"
                        ):
                            st.session_state.pop(
                                f"confirm_delete_trip_{i}", None
                            )
                            st.rerun()


    else:
        with st.container(border=True):
            st.markdown("### 🌤️ Your next adventure starts here")
            st.write(
                "Create a trip above, then add activities "
                "to build your day-by-day itinerary."
            )

    st.write("")
    st.divider()

    # Places You Would Like to Go
    st.markdown(
        '<div class="dashboard-section">'
        '💙 Places You Would Like to Go</div>',
        unsafe_allow_html=True
    )

    st.write(
        "Keep a list of destinations you dream of exploring."
    )

    with st.container(border=True):
        st.markdown("### 🔎 Discover real places")

        with st.form("wishlist_form"):
            place_search = st.text_input(
                "Search for a city, landmark, or attraction",
                placeholder="e.g. Eiffel Tower, Paris or Table Mountain"
            )

            search_places = st.form_submit_button(
                "🔍 Search Places",
                use_container_width=True
            )

        if search_places:
            clean_query = place_search.strip()

            if not clean_query:
                st.warning("Please enter a place to search for.")
                st.session_state.pop("real_place_results", None)
            else:
                try:
                    with st.spinner("Searching real places..."):
                        results = search_real_places(clean_query)

                    st.session_state.real_place_results = results

                    if not results:
                        st.info(
                            "No places found. Try another name or "
                            "include the city and country."
                        )
                except requests.RequestException:
                    st.error(
                        "Place search is temporarily unavailable. "
                        "Please check your connection and try again."
                    )

        results = st.session_state.get("real_place_results", [])

        if results:
            st.write("**Choose a matching place:**")

            # Use the full display name as the label and save its exact coordinates.
            place_options = {
                result["display_name"]: result
                for result in results
                if result.get("display_name")
                and result.get("lat") is not None
                and result.get("lon") is not None
            }

            if place_options:
                selected_place = st.selectbox(
                    "Search results",
                    options=list(place_options.keys()),
                    key="selected_real_place"
                )

                if st.button(
                    "＋ Save Selected Place",
                    key="save_real_place",
                    use_container_width=True
                ):
                    selected_result = place_options[selected_place]
                    clean_place = selected_place.strip()
                    already_saved = any(
                        (
                            item.get("name", "").casefold()
                            if isinstance(item, dict)
                            else str(item).casefold()
                        ) == clean_place.casefold()
                        for item in st.session_state.wishlist
                    )

                    if already_saved:
                        st.warning(
                            "That destination is already in your wishlist."
                        )
                    else:
                        address = selected_result.get("address", {})
                        country = address.get("country", "")
                        st.session_state.wishlist.append({
                            "name": clean_place,
                            "latitude": float(selected_result["lat"]),
                            "longitude": float(selected_result["lon"]),
                            "country": country,
                        })
                        st.success(f"{clean_place} added to your wishlist!")
                        st.rerun()
            else:
                st.info("The search results did not include usable coordinates. Try another search.")

    st.write("")

    if st.session_state.wishlist:
        for i, place in enumerate(st.session_state.wishlist):
            if isinstance(place, dict):
                place_name = place.get("name", "Saved place")
            else:
                place_name = str(place)

            with st.container(border=True):
                place_col, delete_col = st.columns([5, 1])

                with place_col:
                    st.markdown(f"### 🌍 {place_name}")
                    if isinstance(place, dict) and place.get("country"):
                        st.caption(place["country"])

                with delete_col:
                    if st.button(
                        "✕",
                        key=f"delete_place_{i}",
                        use_container_width=True,
                        help="Remove this destination"
                    ):
                        st.session_state[f"confirm_delete_place_{i}"] = True

                if st.session_state.get(f"confirm_delete_place_{i}", False):
                    st.warning(f"Remove {place_name} from your wishlist?")
                    confirm_col, cancel_col = st.columns(2)

                    with confirm_col:
                        if st.button(
                            "Yes, remove",
                            key=f"confirm_delete_place_button_{i}"
                        ):
                            st.session_state.wishlist.pop(i)
                            # Clear index-based confirmations so they cannot
                            # accidentally apply to a different item afterward.
                            for key in list(st.session_state.keys()):
                                if key.startswith("confirm_delete_place_"):
                                    st.session_state.pop(key, None)
                            st.rerun()

                    with cancel_col:
                        if st.button(
                            "Cancel",
                            key=f"cancel_delete_place_{i}"
                        ):
                            st.session_state.pop(
                                f"confirm_delete_place_{i}", None
                            )
                            st.rerun()

    else:
        with st.container(border=True):
            st.write("✨ Your dream destinations will appear here.")
            st.caption(
                "Search for a real place above and save it to your wishlist."
            )
